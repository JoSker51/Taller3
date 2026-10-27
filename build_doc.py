# Genera el documento Word del Taller 3 a partir de los mismos textos y figuras del notebook.
# Requiere haber ejecutado antes el notebook (crea la carpeta figuras/).
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import textos as T

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
TINTA = RGBColor(0x1C, 0x24, 0x33)

def runs(p, text):
    """Agrega texto con **negrita**, *cursiva* y `código`."""
    for tok in re.split(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`)', text):
        if not tok: continue
        if tok.startswith('**'): r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('`'): r = p.add_run(tok[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(10)
        elif tok.startswith('*'): r = p.add_run(tok[1:-1]); r.italic = True
        else: p.add_run(tok)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)

def tabla(lineas):
    filas = [[c.strip() for c in l.strip().strip('|').split('|')] for l in lineas if not re.match(r'^\|\s*-', l)]
    t = doc.add_table(rows=len(filas), cols=len(filas[0])); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, f in enumerate(filas):
        for j, v in enumerate(f):
            c = t.cell(i, j); c.text = ''; p = c.paragraphs[0]; runs(p, v)
            for r in p.runs: r.font.size = Pt(9.5)
            if i == 0:
                shade(c, '1C2433')
                for r in p.runs: r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph()

def md(texto):
    lineas = texto.strip().split('\n'); i = 0
    while i < len(lineas):
        l = lineas[i]
        if l.startswith('|'):
            bloque = []
            while i < len(lineas) and lineas[i].startswith('|'): bloque.append(lineas[i]); i += 1
            tabla(bloque); continue
        if l.startswith('#'):
            nivel = len(l) - len(l.lstrip('#'))
            h = doc.add_heading(l.lstrip('#').strip(), level=min(nivel, 3))
            for r in h.runs: r.font.color.rgb = TINTA
        elif re.match(r'^\s*- ', l):
            p = doc.add_paragraph(style='List Bullet'); runs(p, l.strip()[2:])
        elif re.match(r'^\d+\. ', l):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6); runs(p, l)
        elif l.strip():
            runs(doc.add_paragraph(), l)
        i += 1

def figura(ruta, pie):
    doc.add_picture(ruta, width=Cm(16.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); r = p.add_run(pie); r.italic = True; r.font.size = Pt(9); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Portada
for _ in range(6): doc.add_paragraph()
for txt, tam, neg in [('Taller 3 · Series de tiempo', 26, True), ('Yomi Level: dificultad del japonés en el contenido audiovisual', 15, False),
                      ('', 11, False), ('José Santiago González · Oscar Mateo Arrubla · Leonardo Diasgranados', 12, True),
                      ('Ciencias de la Computación e Inteligencia Artificial · Universidad Sergio Arboleda', 11, False),
                      ('Aprendizaje Supervisado · Docente: Hernán Cruz Bueno', 11, False)]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.font.size = Pt(tam); r.bold = neg; r.font.color.rgb = TINTA
doc.add_page_break()

md(T.INTRO.replace('# Taller 3 · Series de tiempo aplicadas a Yomi Level', '# Contexto'))
md(T.SERIE)
figura('figuras/0_serie.png', 'Figura 0. Dificultad promedio anual (línea) y número de títulos estrenados (barras), 1954-2025.')
md(T.OBS_SERIE)

doc.add_page_break(); md('## 1. Las tres gráficas pedidas')
figura('figuras/1_lag1.png', 'Figura 1. Dificultad real + Lag 1 (año anterior).')
figura('figuras/2_rolling_mean.png', 'Figura 2. Dificultad real + promedio móvil de 3 años (y de 5 años para comparar).')
figura('figuras/3_rolling_std.png', 'Figura 3. Desviación estándar móvil de 3 años.')
figura('figuras/4_combinada.png', 'Figura 4. Las tres variables juntas.')

doc.add_page_break(); md('## 2 y 3. Análisis y respuestas')
for s in (T.RESP_A, T.RESP_B, T.RESP_C, T.RESP_5): md(s)
doc.add_page_break(); md(T.HOJA)

doc.add_page_break(); md(T.MODELOS_INTRO)
figura('figuras/5_comparativa_nivel.png', 'Figura 5. Los 5 modelos prediciendo directamente el nivel (V1), años de prueba 2012-2025.')
md(T.OBS_NIVEL); md(T.OBS_FUGA)
md(T.POR_MODELO_INTRO)
for i, (n, a) in enumerate([('Regresión Lineal', 'regresión'), ('Random Forest', 'random'), ('XGBoost', 'xgboost'), ('SVM', 'svm'), ('KNN', 'knn')]):
    figura(f'figuras/6_{i}_{a}.png', f'Figura {6 + i}. {n}: predicción del nivel (V1) y del cambio (V3) frente a la dificultad real.')
md(T.HALLAZGOS_MODELOS)
doc.add_page_break(); md(T.CONCLUSION)

doc.save('Taller3_Series_Tiempo_Yomi_Level.docx'); print('docx listo')
